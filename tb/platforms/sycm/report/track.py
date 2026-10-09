"""单品诊断报告的执行跟踪（周复盘）：方案开始以后，逐天拿实际的支付金额、转化率和每日门槛比，标出没达标的日子，
出一份填好实际数的执行跟踪表。每周跑一次就是周复盘。只取商品趋势，一两个请求。"""
from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ....core.errors import TbError
from . import collect
from .build import view

RECENT = 7      # 摘要里「最近几天」的日均：小店一天几单，单日达不达标波动太大，按周看


def compare(ramp: list[dict[str, Any]], actual: dict[str, dict[str, Any]], through: date) -> list[dict[str, Any]]:
    days = []
    for r in ramp:
        if r["date"] > through.isoformat():
            break
        a = actual.get(r["date"]) or {}
        amt = a.get("amt")
        days.append({"d": r["d"], "date": r["date"], "phase": r["phase"],
                     "门槛金额": None if r["gmv"] is None else round(r["gmv"], 2), "门槛转化率": r["cv"],
                     "支付金额": amt, "访客数": a.get("uv"), "支付买家数": a.get("byr"), "支付转化率": a.get("cv"),
                     "达标": None if amt is None or r["gmv"] is None else amt >= r["gmv"]})
    return days


def summarize(days: list[dict[str, Any]]) -> dict[str, Any]:
    known = [d for d in days if d["达标"] is not None]
    recent = known[-RECENT:]
    s = lambda rows, k: sum(r[k] or 0 for r in rows)
    streak = 0
    for d in reversed(known):
        if d["达标"]:
            break
        streak += 1
    out: dict[str, Any] = {"已过天数": len(days), "达标天数": sum(1 for d in known if d["达标"]), "有数据的天数": len(known),
                           "连续没达标": streak}
    if recent:
        gmv, need = s(recent, "支付金额") / len(recent), s(recent, "门槛金额") / len(recent)
        uv, byr = s(recent, "访客数"), s(recent, "支付买家数")
        out[f"最近 {len(recent)} 天"] = {"日均支付金额": round(gmv, 2), "日均门槛": round(need, 2),
                                      "差距": f"{(gmv / need - 1) * 100:+.1f}%" if need else "-",
                                      "支付转化率": round(byr / uv, 4) if uv else None,
                                      "转化率门槛（平均）": round(s(recent, "门槛转化率") / len(recent), 4)}
    return out


def run(data: dict[str, Any], analysis: dict[str, Any], *, today: date | None = None) -> dict[str, Any]:
    v = view(data, analysis)
    ramp = v["plan"]["ramp"]
    start, last = date.fromisoformat(ramp[0]["date"]), date.fromisoformat(ramp[-1]["date"])
    through = min((today or date.today()) - timedelta(days=1), last)
    if through < start:
        raise TbError(f"方案还没开始：D1 是 {start}，要等 D1 过完才有数可以对照。", stage="执行跟踪")
    item_id = data["meta"]["item_id"]
    rows = collect.fetch_series(lambda c, d: collect.fetch_item_trend(c, item_id, d), collect.load_taobao_cookies(), through, start)
    days = compare(ramp, {r["d"]: r for r in rows}, through)
    return {"item_id": item_id, "start": start.isoformat(), "through": through.isoformat(), "plan_days": len(ramp),
            "days": days, "summary": summarize(days)}


def render_text(t: dict[str, Any]) -> str:
    s = t["summary"]
    pct = lambda v: "-" if v is None else f"{v * 100:.2f}%"
    lines = [f"# 执行跟踪：商品 {t['item_id']}，方案 D1 = {t['start']}，对照到 {t['through']}（第 {s['已过天数']} 天，共 {t['plan_days']} 天）",
             f"支付金额达标 {s['达标天数']} 天 / 有数据的 {s['有数据的天数']} 天；最近连续没达标 {s['连续没达标']} 天"]
    for k, r in s.items():
        if k.startswith("最近"):
            lines.append(f"{k}：日均支付金额 {r['日均支付金额']}，日均门槛 {r['日均门槛']}（{r['差距']}）；"
                         f"支付转化率 {pct(r['支付转化率'])}，门槛平均 {pct(r['转化率门槛（平均）'])}")
    lines.append("\n天\t日期\t门槛金额\t实际金额\t门槛转化率\t实际转化率\t达标")
    for d in t["days"]:
        mark = "-" if d["达标"] is None else ("达标" if d["达标"] else "没达标")
        lines.append(f"D{d['d']}\t{d['date']}\t{d['门槛金额']}\t{d['支付金额'] if d['支付金额'] is not None else '-'}\t"
                     f"{pct(d['门槛转化率'])}\t{pct(d['支付转化率'])}\t{mark}")
    return "\n".join(lines)


def write_xlsx(t: dict[str, Any], path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "执行跟踪"
    ws.append(["天", "日期", "阶段", "日均支付金额门槛", "实际支付金额", "支付转化率门槛", "实际支付转化率", "访客数", "支付买家数", "达标"])
    for c in ws[1]:
        c.font, c.fill, c.alignment = Font(bold=True), PatternFill("solid", fgColor="EAF1FE"), Alignment(horizontal="center", wrap_text=True)
    red = PatternFill("solid", fgColor="FDEDED")
    for d in t["days"]:
        ws.append([f"D{d['d']}", d["date"], d["phase"], d["门槛金额"], d["支付金额"], d["门槛转化率"], d["支付转化率"],
                   d["访客数"], d["支付买家数"], "-" if d["达标"] is None else ("达标" if d["达标"] else "没达标")])
        if d["达标"] is False:
            ws.cell(ws.max_row, 10).fill = red
    for row in ws.iter_rows(min_row=2):
        row[3].number_format = row[4].number_format = "#,##0.00"
        row[5].number_format = row[6].number_format = "0.00%"
    for col, w in zip("ABCDEFGHIJ", (6, 12, 8, 16, 14, 16, 16, 10, 12, 8)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(path)
    return path
