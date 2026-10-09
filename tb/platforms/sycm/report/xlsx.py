"""单品诊断报告的配套表格：执行跟踪（每日门槛，实际数留空自己填，表里自动算达不达标）、动作清单、每日数据、两个周期对比、
同店对标、流量来源、SKU。用和网页同一份数据（脱敏版也一样脱敏）。"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="EAF1FE")
FILL_ME = PatternFill("solid", fgColor="FFF4E5")       # 要自己填的列
PCT, MONEY, NUM = "0.00%", "#,##0.00", "#,##0.0"


def _sheet(wb: Workbook, title: str, head: list[str], rows: list[list[Any]], formats: dict[int, str] | None = None,
           widths: dict[int, int] | None = None):
    ws = wb.create_sheet(title)
    ws.append(head)
    for c in ws[1]:
        c.font, c.fill, c.alignment = Font(bold=True), HEAD_FILL, Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for col, fmt in (formats or {}).items():
        for cell in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cell[0].number_format = fmt
    for i in range(1, len(head) + 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, max(10, len(head[i - 1]) * 2 + 2))
    ws.freeze_panes = "A2"
    return ws


def write(v: dict[str, Any], path: Path) -> Path:
    unit = "（指数）" if v["masked"] else ""
    wb = Workbook()
    wb.remove(wb.active)
    analysis, plan = v["analysis"], v["plan"]
    names = {n: p.get("name") or f"阶段{n}" for n, p in enumerate(analysis["plan"]["phases"], 1)}

    # 执行跟踪：门槛是程序算的，F～H 自己每天填，I、J 自动算
    rows = []
    for i, r in enumerate(plan["ramp"], 2):
        rows.append([f"D{r['d']}", r["date"], names.get(r["phase"], ""), r["gmv"], r["cv"], None, None, None,
                     f'=IF(AND(G{i}<>"",G{i}>0),H{i}/G{i},"")', f'=IF(F{i}="","",IF(F{i}>=D{i},"达标","没达标"))'])
    ws = _sheet(wb, "执行跟踪", ["天", "日期", "阶段", f"日均支付金额门槛{unit}", "支付转化率门槛", f"实际支付金额{unit}（填）",
                                "实际访客（填）", "实际支付买家（填）", "实际转化率", "支付金额达标"],
                rows, {4: MONEY, 5: PCT, 6: MONEY, 9: PCT}, {3: 16, 4: 18, 6: 18})
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=8):
        for cell in row:
            cell.fill = FILL_ME

    _sheet(wb, "动作清单", ["#", "阶段", "动作", "负责", "看什么指标", "时间", "状态（填）", "备注（填）"],
           [[n, names.get(a.get("phase"), ""), a.get("what"), a.get("owner"), a.get("watch"), a.get("when"), None, None]
            for n, a in enumerate(analysis["plan"].get("actions") or [], 1)], widths={3: 60, 5: 22, 6: 18})

    cols = [("d", "日期", None), ("uv", "访客数", NUM), ("pv", "浏览量", NUM), ("amt", f"支付金额{unit}", MONEY),
            ("itm", "支付件数", NUM), ("byr", "支付买家数", NUM), ("cv", "支付转化率", PCT), ("aov", f"客单价{unit}", MONEY),
            ("cart", "加购件数", NUM), ("cartRate", "加购率", PCT), ("clt", "收藏人数", NUM), ("newb", "新买家", NUM),
            ("oldb", "老买家", NUM), ("bounce", "跳失率", PCT), ("stay", "停留（秒）", NUM), ("charge", f"推广花费{unit}", MONEY),
            ("rfd", f"退款金额{unit}", MONEY)]
    search = [("seUv", "搜索引导访客", NUM), ("seShare", "搜索引导访客占比", PCT)]
    _sheet(wb, "每日数据", [c[1] for c in cols + search], [[r.get(k) for k, _, _ in cols + search] for r in v["daily"]],
           {i: f for i, (_, _, f) in enumerate(cols + search, 1) if f})

    h0, h1 = v["halves"]["对比周期"], v["halves"]["分析周期"]
    _sheet(wb, "两个周期对比", ["指标", "对比周期", "分析周期"],
           [[name, h0.get(k), h1.get(k)] for k, name, _ in cols[1:]], widths={1: 16})

    _sheet(wb, "同店对标", ["商品", "类目", "同类目", "支付金额" + unit, "访客数", "支付转化率", "加购率（加购件数÷访客）", "客单价" + unit,
                         "搜索引导访客", "跳失率"],
           [[r.get("title") or r.get("id"), r.get("cateName") or r.get("cate"), "本品" if r.get("me") else ("是" if r.get("same") else ""),
             r.get("amt"), r.get("uv"), r.get("cv"), r.get("cartPerUv"), r.get("aov"), r.get("seUv"), r.get("bounce")] for r in v["bench"]],
           {4: MONEY, 5: NUM, 6: PCT, 7: PCT, 8: MONEY, 9: NUM, 10: PCT}, {1: 40, 2: 16})

    flows = {p: {r["path"]: r for r in rows} for p, rows in v["flow"].items()}
    paths = list(dict.fromkeys([*flows.get("分析周期", {}), *flows.get("对比周期", {})]))
    _sheet(wb, "流量来源", ["渠道", "访客（对比）", "访客（分析）", "支付买家（对比）", "支付买家（分析）", "支付转化率（对比）", "支付转化率（分析）"],
           [[p, flows.get("对比周期", {}).get(p, {}).get("uv"), flows.get("分析周期", {}).get(p, {}).get("uv"),
             flows.get("对比周期", {}).get(p, {}).get("byr"), flows.get("分析周期", {}).get(p, {}).get("byr"),
             flows.get("对比周期", {}).get(p, {}).get("cv"), flows.get("分析周期", {}).get(p, {}).get("cv")] for p in paths],
           {2: NUM, 3: NUM, 4: NUM, 5: NUM, 6: PCT, 7: PCT}, {1: 30})

    sku = (v.get("deep") or {}).get("sku") or []
    if sku:
        _sheet(wb, "SKU", ["SKU", "加购件数", "支付金额" + unit, "支付件数", "支付买家数"],
               [[r["sku"], r["cart"], r["amt"], r["itm"], r["byr"]] for r in sku], {2: NUM, 3: MONEY, 4: NUM, 5: NUM}, {1: 44})
    wb.save(path)
    return path
